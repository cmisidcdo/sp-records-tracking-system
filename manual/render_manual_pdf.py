from __future__ import annotations

import html
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pypdfium2 import PdfDocument
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    KeepTogether,
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
DOCX_PATH = BASE / "Sangguniang_Panlungsod_Records_Tracking_System_User_Manual.docx"
PDF_PATH = BASE / "Sangguniang_Panlungsod_Records_Tracking_System_User_Manual.pdf"
RENDER_DIR = BASE / "rendered"
LOGO_PATH = ROOT / "public" / "assets" / "splogo.jpg"

NAVY = colors.HexColor("#173A5E")
BLUE = colors.HexColor("#1E6FA8")
LIGHT_BLUE = colors.HexColor("#DDEBF7")
PALE_BLUE = colors.HexColor("#EEF5FA")
GOLD = colors.HexColor("#F2CF32")
PALE_GOLD = colors.HexColor("#FFF6CC")
GREEN = colors.HexColor("#2E7D5A")
PALE_GREEN = colors.HexColor("#E7F4EC")
RED = colors.HexColor("#B42318")
PALE_RED = colors.HexColor("#FDECEC")
INK = colors.HexColor("#263746")
MID = colors.HexColor("#52697D")
LINE = colors.HexColor("#B9C8D4")
GRAY = colors.HexColor("#F3F6F8")
WHITE = colors.white


class ManualDocTemplate(SimpleDocTemplate):
    def __init__(self, filename, **kwargs):
        super().__init__(filename, **kwargs)
        self.heading_count = 0

    def beforeDocument(self):
        self.heading_count = 0

    def afterFlowable(self, flowable):
        if not isinstance(flowable, Paragraph):
            return
        name = flowable.style.name
        if name not in {"Heading1", "Heading2", "Heading3"}:
            return
        level = {"Heading1": 0, "Heading2": 1, "Heading3": 2}[name]
        text = flowable.getPlainText()
        self.heading_count += 1
        key = f"h{self.heading_count}"
        self.canv.bookmarkPage(key)
        self.canv.addOutlineEntry(text, key, level=level, closed=False)


def iter_block_items(parent):
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, parent)


def has_page_break(paragraph: DocxParagraph) -> bool:
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def markup_for_paragraph(paragraph: DocxParagraph) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        text = html.escape(run.text).replace("\n", "<br/>")
        if not text:
            continue
        if run.bold:
            text = f"<b>{text}</b>"
        if run.italic:
            text = f"<i>{text}</i>"
        parts.append(text)
    return "".join(parts) or html.escape(paragraph.text)


def build_styles():
    sample = getSampleStyleSheet()
    styles = {
        "Body": ParagraphStyle(
            "Body",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=9.1,
            leading=10.8,
            textColor=INK,
            spaceAfter=4.2,
            allowWidows=0,
            allowOrphans=0,
        ),
        "BodyCenter": ParagraphStyle(
            "BodyCenter",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=9.1,
            leading=10.8,
            textColor=INK,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
        "Heading1": ParagraphStyle(
            "Heading1",
            parent=sample["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=19,
            leading=21,
            textColor=NAVY,
            spaceBefore=8,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "Heading2": ParagraphStyle(
            "Heading2",
            parent=sample["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12.2,
            leading=14,
            textColor=BLUE,
            spaceBefore=7,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "Heading3": ParagraphStyle(
            "Heading3",
            parent=sample["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=10.1,
            leading=11.5,
            textColor=NAVY,
            spaceBefore=5,
            spaceAfter=3,
            keepWithNext=True,
        ),
        "Bullet": ParagraphStyle(
            "Bullet",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=8.9,
            leading=10.5,
            textColor=INK,
            leftIndent=17,
            firstLineIndent=-10,
            bulletIndent=4,
            spaceAfter=2.5,
        ),
        "Number": ParagraphStyle(
            "Number",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=8.9,
            leading=10.5,
            textColor=INK,
            leftIndent=19,
            firstLineIndent=-14,
            bulletIndent=2,
            spaceAfter=3,
        ),
        "Table": ParagraphStyle(
            "Table",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=8.8,
            textColor=INK,
            spaceAfter=0,
        ),
        "TableHeader": ParagraphStyle(
            "TableHeader",
            parent=sample["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=7.8,
            leading=9,
            textColor=WHITE,
            spaceAfter=0,
        ),
        "CalloutTitle": ParagraphStyle(
            "CalloutTitle",
            parent=sample["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8.6,
            leading=9.5,
            textColor=BLUE,
            spaceAfter=2,
        ),
    }
    return styles


STYLES = build_styles()


def draw_later_page(canvas, doc):
    if doc.page <= 1:
        return
    canvas.saveState()
    canvas.setStrokeColor(GOLD)
    canvas.setLineWidth(1.1)
    canvas.line(0.78 * inch, 10.38 * inch, 7.72 * inch, 10.38 * inch)
    canvas.setFont("Helvetica-Bold", 7.2)
    canvas.setFillColor(NAVY)
    canvas.drawString(0.78 * inch, 10.48 * inch, "SP RECORDS TRACKING SYSTEM  |  USER MANUAL")
    canvas.setFont("Helvetica-Bold", 6.8)
    canvas.setFillColor(MID)
    canvas.drawRightString(7.72 * inch, 10.48 * inch, "OFFICE OF THE SANGGUNIANG PANLUNGSOD")
    canvas.setStrokeColor(LINE)
    canvas.setLineWidth(0.45)
    canvas.line(0.78 * inch, 0.48 * inch, 7.72 * inch, 0.48 * inch)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(MID)
    canvas.drawCentredString(
        4.25 * inch,
        0.31 * inch,
        f"Internal User Guide  |  Version 1.0  |  July 2026  |  Page {doc.page}",
    )
    canvas.restoreState()


def build_cover(story):
    story.extend([Spacer(1, 0.46 * inch)])
    story.append(Image(str(LOGO_PATH), width=1.55 * inch, height=1.55 * inch))
    for text, size, color, leading, after in [
        ("REPUBLIC OF THE PHILIPPINES", 8.5, MID, 10, 0),
        ("CAGAYAN DE ORO CITY", 8.5, MID, 10, 0),
        ("OFFICE OF THE SANGGUNIANG PANLUNGSOD", 10.5, NAVY, 12, 16),
        ("SANGGUNIANG PANLUNGSOD", 20, NAVY, 21, 2),
        ("RECORDS TRACKING SYSTEM", 23, BLUE, 24, 12),
        ("USER'S MANUAL", 16, INK, 18, 18),
    ]:
        story.append(
            Paragraph(
                f"<b>{text}</b>",
                ParagraphStyle(
                    f"Cover-{text}",
                    fontName="Helvetica-Bold",
                    fontSize=size,
                    leading=leading,
                    textColor=color,
                    alignment=TA_CENTER,
                    spaceAfter=after,
                ),
            )
        )
    meta = Table(
        [
            [
                Paragraph("<b>VERSION</b><br/>1.0", STYLES["BodyCenter"]),
                Paragraph("<b>RELEASE</b><br/>July 2026", STYLES["BodyCenter"]),
                Paragraph("<b>AUDIENCE</b><br/>Authorized Users", STYLES["BodyCenter"]),
            ]
        ],
        colWidths=[2.29 * inch] * 3,
    )
    meta.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), LIGHT_BLUE),
                ("BOX", (0, 0), (-1, -1), 0.5, WHITE),
                ("INNERGRID", (0, 0), (-1, -1), 1, WHITE),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(meta)
    story.append(Spacer(1, 0.24 * inch))
    story.append(
        Paragraph(
            "Operational guide for records intake, review, committee action,<br/>"
            "plenary processing, post-plenary tracking, reporting, and administration",
            ParagraphStyle(
                "CoverFooter",
                fontName="Helvetica",
                fontSize=9.2,
                leading=11,
                textColor=MID,
                alignment=TA_CENTER,
            ),
        )
    )
    story.append(PageBreak())


def docx_table_to_reportlab(block: DocxTable):
    data = []
    for row_index, row in enumerate(block.rows):
        rendered_row = []
        for cell in row.cells:
            paragraphs = []
            for p in cell.paragraphs:
                text = markup_for_paragraph(p)
                if text:
                    paragraphs.append(text)
            combined = "<br/>".join(paragraphs)
            style = STYLES["TableHeader"] if row_index == 0 and len(block.rows) > 1 else STYLES["Table"]
            rendered_row.append(Paragraph(combined or " ", style))
        data.append(rendered_row)

    cols = max(len(row) for row in data)
    raw_widths = []
    for idx in range(cols):
        width = block.rows[0].cells[idx].width
        raw_widths.append(float(width or 1))
    minimum_width = 0.3 * inch
    raw_widths = [max(width, minimum_width) for width in raw_widths]
    total = sum(raw_widths) or cols
    max_width = 6.87 * inch
    col_widths = [max_width * width / total for width in raw_widths]

    table = LongTable(data, colWidths=col_widths, repeatRows=1 if len(data) > 1 else 0, hAlign="CENTER")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.35, LINE),
    ]
    if len(data) > 1:
        commands.append(("BACKGROUND", (0, 0), (-1, 0), NAVY))
        for row_idx in range(1, len(data)):
            commands.append(("BACKGROUND", (0, row_idx), (-1, row_idx), GRAY if row_idx % 2 == 0 else WHITE))
    else:
        commands.append(("BACKGROUND", (0, 0), (-1, -1), PALE_BLUE))
    table.setStyle(TableStyle(commands))
    return table


def build_story_from_docx():
    doc = Document(DOCX_PATH)
    story = []
    build_cover(story)
    seen_first_break = False
    list_number = 0
    toc_seen = False

    for block in iter_block_items(doc):
        if not seen_first_break:
            if isinstance(block, DocxParagraph) and has_page_break(block):
                seen_first_break = True
            continue

        if isinstance(block, DocxTable):
            story.append(docx_table_to_reportlab(block))
            story.append(Spacer(1, 5))
            list_number = 0
            continue

        paragraph = block
        if has_page_break(paragraph):
            story.append(PageBreak())
            story.append(Spacer(1, (0.55 if toc_seen else 0.15) * inch))
            toc_seen = False
            list_number = 0
            continue
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else "Normal"

        if "Right-click and select Update Field" in text:
            entries = [
                ("Document Control", "2"),
                ("Contents", "3"),
                ("About the System", "4"),
                ("Roles and Responsibilities", "5"),
                ("Navigation and Dashboards", "6"),
                ("Create a New Record", "7"),
                ("Committee Referral Workflow", "8"),
                ("Update Status and History", "9"),
                ("Plenary and Post-Plenary Processing", "11"),
                ("Transmittals, Letters and Endorsements", "13"),
                ("Records, Search, Attachments, and History", "14"),
                ("Committee Referral Printing", "15"),
                ("Terms, Officials, Committees, and Assignments", "16"),
                ("User Accounts", "17"),
                ("Reports, Audit Logs, and Backup", "18"),
                ("Running on a Local Network", "19"),
                ("Troubleshooting", "20"),
                ("Quick Reference", "21"),
            ]
            toc_data = [
                [
                    Paragraph(f"<b>{html.escape(title)}</b>", STYLES["Body"]),
                    Paragraph(f"<b>{page}</b>", STYLES["BodyCenter"]),
                ]
                for title, page in entries
            ]
            toc = Table(toc_data, colWidths=[6.25 * inch, 0.62 * inch], hAlign="CENTER")
            toc.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                    ]
                )
            )
            story.append(toc)
            toc_seen = True
            continue

        markup = markup_for_paragraph(paragraph)
        if style_name.startswith("Heading 1"):
            story.append(Paragraph(markup, STYLES["Heading1"]))
            list_number = 0
        elif style_name.startswith("Heading 2"):
            story.append(Paragraph(markup, STYLES["Heading2"]))
            list_number = 0
        elif style_name.startswith("Heading 3"):
            story.append(Paragraph(markup, STYLES["Heading3"]))
            list_number = 0
        elif style_name.startswith("List Bullet"):
            story.append(Paragraph(markup, STYLES["Bullet"], bulletText="•"))
            list_number = 0
        elif style_name.startswith("List Number"):
            list_number += 1
            story.append(Paragraph(markup, STYLES["Number"], bulletText=f"{list_number}."))
        else:
            if not text.startswith(r"C:\dbase\php\php.exe"):
                list_number = 0
            alignment = STYLES["BodyCenter"] if paragraph.alignment == 1 else STYLES["Body"]
            story.append(Paragraph(markup, alignment))
    return story


def build_pdf():
    template = ManualDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=0.78 * inch,
        rightMargin=0.78 * inch,
        topMargin=0.80 * inch,
        bottomMargin=0.70 * inch,
        title="Sangguniang Panlungsod Records Tracking System User's Manual",
        author="Office of the Sangguniang Panlungsod",
    )
    template.build(
        build_story_from_docx(),
        onFirstPage=lambda canvas, doc: None,
        onLaterPages=draw_later_page,
    )


def render_pdf():
    RENDER_DIR.mkdir(parents=True, exist_ok=True)
    for old in RENDER_DIR.glob("page-*.png"):
        old.unlink()
    pdf = PdfDocument(str(PDF_PATH))
    scale = 130 / 72
    for index, page in enumerate(pdf):
        bitmap = page.render(scale=scale)
        image = bitmap.to_pil()
        image.save(RENDER_DIR / f"page-{index + 1:03d}.png")
    print(f"PDF={PDF_PATH}")
    print(f"PAGES={len(pdf)}")
    print(f"RENDER_DIR={RENDER_DIR}")


if __name__ == "__main__":
    build_pdf()
    render_pdf()
