from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = Path(__file__).resolve().parent
LOGO_PATH = ROOT / "public" / "assets" / "splogo.jpg"
OUTPUT_PATH = OUTPUT_DIR / "Sangguniang_Panlungsod_Records_Tracking_System_User_Manual.docx"

NAVY = "173A5E"
BLUE = "1E6FA8"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "EEF5FA"
GOLD = "F2CF32"
PALE_GOLD = "FFF6CC"
GREEN = "2E7D5A"
PALE_GREEN = "E7F4EC"
RED = "B42318"
PALE_RED = "FDECEC"
ORANGE = "C46B18"
PALE_ORANGE = "FFF0DF"
INK = "263746"
MID = "52697D"
LINE = "B9C8D4"
WHITE = "FFFFFF"
GRAY = "F3F6F8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in edges[edge].items():
            element.set(qn("w:" + key), str(value))


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_grid(table, widths: Sequence[float]) -> None:
    table.autofit = False
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(Inches(width).emu / 635)))
        tbl_grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])


def set_paragraph_border(paragraph, color=LINE, size=8, space=3, side="bottom") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-1" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep])
    placeholder = paragraph.add_run("Right-click and select Update Field to refresh this contents list.")
    placeholder.italic = True
    placeholder.font.color.rgb = RGBColor.from_string(MID)
    paragraph.add_run()._r.append(fld_end)


def add_external_hyperlink(paragraph, text: str, url: str, color=BLUE) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_internal_hyperlink(paragraph, text: str, anchor: str, color=BLUE) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def set_run_font(run, size=None, bold=None, color=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def style_paragraph(paragraph, space_before=0, space_after=4, line=1.05) -> None:
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None, italic=False):
    p = doc.add_paragraph()
    style_paragraph(p, space_after=5, line=1.08)
    if bold_prefix and text.startswith(bold_prefix):
        lead = p.add_run(bold_prefix)
        set_run_font(lead, bold=True, color=INK)
        tail = p.add_run(text[len(bold_prefix):])
        set_run_font(tail, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, color=INK, italic=italic)
    return p


def add_bullet(doc: Document, text: str, *, level=0, bold_prefix: str | None = None):
    style_name = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.left_indent = Inches(0.23 + level * 0.22)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    style_paragraph(p, space_after=2.5, line=1.04)
    if bold_prefix and text.startswith(bold_prefix):
        lead = p.add_run(bold_prefix)
        set_run_font(lead, bold=True, color=INK)
        tail = p.add_run(text[len(bold_prefix):])
        set_run_font(tail, color=INK)
    else:
        set_run_font(p.add_run(text), color=INK)
    return p


def add_number(doc: Document, text: str, *, level=0, bold_prefix: str | None = None):
    style_name = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.22)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    style_paragraph(p, space_after=3, line=1.05)
    if bold_prefix and text.startswith(bold_prefix):
        lead = p.add_run(bold_prefix)
        set_run_font(lead, bold=True, color=INK)
        tail = p.add_run(text[len(bold_prefix):])
        set_run_font(tail, color=INK)
    else:
        set_run_font(p.add_run(text), color=INK)
    return p


def add_heading(doc: Document, text: str, level=1, bookmark: str | None = None, bookmark_id=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    if bookmark:
        add_bookmark(p, bookmark, bookmark_id)
    return p


def add_section_page(doc: Document, section_no: str, title: str, subtitle: str):
    doc.add_page_break()
    p = doc.add_paragraph()
    style_paragraph(p, space_after=2)
    run = p.add_run(section_no.upper())
    set_run_font(run, size=9, bold=True, color=BLUE)
    set_paragraph_border(p, color=GOLD, size=16, space=4, side="bottom")
    h = add_heading(doc, title, 1)
    h.paragraph_format.space_before = Pt(7)
    h.paragraph_format.space_after = Pt(4)
    s = doc.add_paragraph()
    style_paragraph(s, space_after=12, line=1.05)
    set_run_font(s.add_run(subtitle), size=10.5, color=MID, italic=True)


def add_callout(doc: Document, title: str, text: str, kind="info"):
    colors = {
        "info": (PALE_BLUE, BLUE),
        "important": (PALE_GOLD, NAVY),
        "warning": (PALE_RED, RED),
        "success": (PALE_GREEN, GREEN),
    }
    fill, accent = colors[kind]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_grid(table, [0.09, 6.78])
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 0), 0, 0, 0, 0)
    set_cell_margins(table.cell(0, 1), 110, 150, 110, 150)
    table.cell(0, 0).text = ""
    p = table.cell(0, 1).paragraphs[0]
    style_paragraph(p, space_after=2, line=1.03)
    set_run_font(p.add_run(title), size=9.5, bold=True, color=accent)
    p2 = table.cell(0, 1).add_paragraph()
    style_paragraph(p2, space_after=0, line=1.04)
    set_run_font(p2.add_run(text), size=9.4, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_kpi_strip(doc: Document, items: Sequence[tuple[str, str]], fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [6.87 / len(items)] * len(items)
    set_table_grid(table, widths)
    for index, (title, detail) in enumerate(items):
        cell = table.cell(0, index)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, 100, 120, 100, 120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=1, line=1)
        set_run_font(p.add_run(title), size=9, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p2, space_after=0, line=1)
        set_run_font(p2.add_run(detail), size=8.6, color=INK)
        border = {"val": "single", "sz": "6", "color": WHITE}
        set_cell_border(cell, top=border, left=border, bottom=border, right=border)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths: Sequence[float],
    *,
    font_size=8.7,
    header_fill=NAVY,
    alternate=True,
):
    data = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_grid(table, widths)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, 90, 95, 90, 95)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(p, space_after=0, line=1)
        set_run_font(p.add_run(header), size=8.8, bold=True, color=WHITE)
    for row_index, values in enumerate(data):
        row = table.add_row()
        prevent_row_split(row)
        if alternate and row_index % 2:
            fill = GRAY
        else:
            fill = WHITE
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            set_cell_shading(cell, fill)
            set_cell_margins(cell, 80, 95, 80, 95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            style_paragraph(p, space_after=0, line=1.02)
            set_run_font(p.add_run(str(value)), size=font_size, color=INK)
            border = {"val": "single", "sz": "4", "color": LINE}
            set_cell_border(cell, top=border, left=border, bottom=border, right=border)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_workflow(doc: Document, stages: Sequence[tuple[str, str, str]]):
    table = doc.add_table(rows=1, cols=len(stages))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [6.87 / len(stages)] * len(stages)
    set_table_grid(table, widths)
    fills = [PALE_BLUE, PALE_GOLD, PALE_GREEN, PALE_ORANGE, GRAY, PALE_BLUE]
    for index, (number, owner, action) in enumerate(stages):
        cell = table.cell(0, index)
        set_cell_shading(cell, fills[index % len(fills)])
        set_cell_margins(cell, 100, 90, 100, 90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=2, line=1)
        set_run_font(p.add_run(number), size=13, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p2, space_after=2, line=1)
        set_run_font(p2.add_run(owner), size=8.5, bold=True, color=INK)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p3, space_after=0, line=1)
        set_run_font(p3.add_run(action), size=8, color=INK)
        border = {"val": "single", "sz": "6", "color": WHITE}
        set_cell_border(cell, top=border, left=border, bottom=border, right=border)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def format_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.06

    for style_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor.from_string(INK)

    heading_specs = {
        "Heading 1": (20, NAVY, 12, 7),
        "Heading 2": (13.5, BLUE, 9, 4),
        "Heading 3": (10.5, NAVY, 6, 2),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Button" not in styles:
        button = styles.add_style("Button", WD_STYLE_TYPE.CHARACTER)
        button.font.name = "Calibri"
        button._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        button.font.bold = True
        button.font.color.rgb = RGBColor.from_string(BLUE)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(p, space_after=0, line=1)
    run = p.add_run("SP RECORDS TRACKING SYSTEM  |  USER MANUAL")
    set_run_font(run, size=8, bold=True, color=NAVY)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(4.35))
    p.add_run("\t")
    run = p.add_run("OFFICE OF THE SANGGUNIANG PANLUNGSOD")
    set_run_font(run, size=7.5, bold=True, color=MID)
    set_paragraph_border(p, color=GOLD, size=12, space=3, side="bottom")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(fp, space_after=0, line=1)
    set_paragraph_border(fp, color=LINE, size=4, space=3, side="top")
    set_run_font(fp.add_run("Internal User Guide  |  Version 1.0  |  July 2026  |  Page "), size=7.5, color=MID)
    add_page_number(fp)

    core = doc.core_properties
    core.title = "Sangguniang Panlungsod Records Tracking System User Manual"
    core.subject = "Role-based operating guide for the SP Records Tracking System"
    core.author = "Office of the Sangguniang Panlungsod"
    core.keywords = "records tracking, committee referrals, user manual, Sangguniang Panlungsod"
    core.comments = "Version 1.0 - July 2026"


def build_cover(doc: Document) -> None:
    for _ in range(2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.55))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=1, line=1)
    set_run_font(p.add_run("REPUBLIC OF THE PHILIPPINES"), size=9, bold=True, color=MID)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=1, line=1)
    set_run_font(p.add_run("CAGAYAN DE ORO CITY"), size=9, bold=True, color=MID)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=18, line=1)
    set_run_font(p.add_run("OFFICE OF THE SANGGUNIANG PANLUNGSOD"), size=11, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=2, line=0.95)
    set_run_font(p.add_run("SANGGUNIANG PANLUNGSOD"), size=22, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=10, line=0.95)
    set_run_font(p.add_run("RECORDS TRACKING SYSTEM"), size=25, bold=True, color=BLUE)
    set_paragraph_border(p, color=GOLD, size=22, space=9, side="bottom")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=8, space_after=18, line=1)
    set_run_font(p.add_run("USER'S MANUAL"), size=17, bold=True, color=INK)

    add_kpi_strip(
        doc,
        [
            ("VERSION", "1.0"),
            ("RELEASE", "July 2026"),
            ("AUDIENCE", "Authorized Users"),
        ],
        fill=PALE_BLUE,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("Operational guide for records intake, review, committee action,"), size=9.5, color=MID)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=0, line=1)
    set_run_font(p.add_run("plenary processing, post-plenary tracking, reporting, and administration"), size=9.5, color=MID)


def build_front_matter(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Document Control", 1)
    add_table(
        doc,
        ["Item", "Details"],
        [
            ("Document title", "Sangguniang Panlungsod Records Tracking System User's Manual"),
            ("Version", "1.0"),
            ("Issue date", "July 2026"),
            ("Owner", "Office of the Sangguniang Panlungsod"),
            ("Coverage", "Current PHP/MySQL development system and authorized network use"),
            ("Classification", "Internal operational guide"),
        ],
        [1.65, 5.22],
        font_size=9.2,
    )

    add_heading(doc, "Revision Record", 2)
    add_table(
        doc,
        ["Version", "Date", "Summary"],
        [("1.0", "July 2026", "Initial manual based on the current configured system.")],
        [0.9, 1.3, 4.67],
        font_size=9.2,
    )

    add_callout(
        doc,
        "Important",
        "The functions visible to a user depend on the account role, assigned Division, and assigned committee. Screens may therefore differ from the examples and instructions in this manual.",
        "important",
    )

    doc.add_page_break()
    add_heading(doc, "Contents", 1)
    add_table(
        doc,
        ["Section", "Page"],
        [
            ("Document Control", "2"),
            ("About the System", "4"),
            ("Roles and Responsibilities", "5"),
            ("Navigation and Dashboards", "6"),
            ("Create a New Record", "7"),
            ("Committee Referral Workflow", "8"),
            ("Update Status and History", "9"),
            ("Plenary and Post-Plenary Processing", "11"),
            ("Transmittals, Letters and Endorsements", "13"),
            ("Records, Search, Attachments, and History", "14"),
            ("Committee Referral Printing", "15"),
            ("Terms, Officials, Committees, and Assignments", "16"),
            ("User Accounts", "17"),
            ("Reports, Audit Logs, and Backup", "18"),
            ("Running on a Local Network", "19"),
            ("Troubleshooting", "20"),
            ("Quick Reference", "21"),
        ],
        [6.12, 0.75],
        font_size=9.2,
    )


def build_introduction(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 1",
        "About the System",
        "Purpose, terminology, access, and the basic operating rules that apply to every user.",
    )
    add_heading(doc, "1.1 Purpose", 2)
    add_body(
        doc,
        "The Sangguniang Panlungsod Records Tracking System records the complete life cycle of Committee Referrals and Transmittals, Letters and Endorsements. It provides a shared source of truth for intake, review, committee assignment, staff action, plenary processing, post-plenary transmittal, reporting, and audit.",
    )
    add_body(
        doc,
        "The system is designed for concurrent use. Every user must sign in with an individual account so actions, remarks, attachments, and status changes are attributed correctly.",
    )

    add_heading(doc, "1.2 Core Terms", 2)
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ("Communication Number", "The unique public tracking number. Committee Referrals use L-00001-YYYY; Transmittals, Letters and Endorsements use A-00001-YYYY."),
            ("Committee Referral", "A communication reviewed and referred to one or more legislative committees."),
            ("Lead Committee", "The primary committee. It is automatic for one committee and selected when several committees are assigned."),
            ("Movement / Update", "A dated status action recorded in the Status Tracking History."),
            ("Action Required", "A record waiting for the signed-in user to review, print, update, or complete an assigned task."),
            ("Approved in the Plenary", "A locked legislative result with an Ordinance or Resolution number and approval date."),
        ],
        [1.55, 5.32],
    )

    add_heading(doc, "1.3 Access Addresses", 2)
    add_bullet(doc, "On the host computer: open http://127.0.0.1:8080/ or the address configured by the Administrator.")
    add_bullet(doc, "On another computer: open http://HOST-IP:8080/. Replace HOST-IP with the local network address of the host computer.")
    add_bullet(doc, "Do not use 127.0.0.1 on another computer. That address always points back to the computer currently being used.")

    add_callout(
        doc,
        "One server, one database",
        "All users must open the same host address. Running separate copies on several computers creates separate databases and changes will not synchronize.",
        "warning",
    )

    add_heading(doc, "1.4 Sign In and Sign Out", 2)
    add_number(doc, "Open the system address.")
    add_number(doc, "Select Sign In, then enter the email address and password assigned to you.")
    add_number(doc, "Select Sign In. The system opens the dashboard allowed for your role.")
    add_number(doc, "When finished, open Profile and select Sign Out.")
    add_callout(
        doc,
        "Account security",
        "Do not share accounts or passwords. Passwords are stored securely and cannot be displayed after creation. An Administrator or City Secretary can reset a password when necessary.",
        "important",
    )


def build_roles(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 2",
        "Roles and Responsibilities",
        "Use this section to understand what each account may see, create, update, approve, or administer.",
    )
    add_heading(doc, "2.1 Role Matrix", 2)
    add_table(
        doc,
        ["Role", "Primary responsibility", "Main access", "Important restriction"],
        [
            ("Administrator", "System-wide administration and oversight.", "All records, dashboards, management, users, audit logs, backup, reports, and role-specific functions.", "Use elevated access only for authorized work."),
            ("City Secretary", "Review and final classification; committee and plenary action.", "All records; review/assignment; terms, officials, committees, users; plenary numbers; reports.", "Ordinary committee status updates are handed to the Division Chief and Secretariat after assignment, except plenary functions."),
            ("Receiving Section Staff", "Initial receipt, data entry, attachment intake, and initial referral printing.", "Create records, search/view, edit or delete own unassigned Received records, print and mark forwarded.", "Cannot update committee status."),
            ("Division Chief", "First committee action, oversight, and review of Secretariat work.", "Assigned committee records, first action, status updates, comments, reports, committee/staff dashboards.", "Limited to assigned committees and staff."),
            ("Secretariat", "Committee processing and status updates.", "Assigned committee records, attachments, history, status updates after Chief action, update referral printing.", "Limited to assigned committees; cannot act before the Chief's first action."),
            ("Division Staff", "Search and read-only support.", "View records, full status history, and reports within the assigned Division scope.", "Cannot create, edit, delete, or update records."),
            ("LMIS & Records Staff", "Post-plenary processing and transmittal.", "Approved records, post-plenary statuses, recipients, history, and permitted contact details.", "Does not manage committee-stage actions."),
        ],
        [1.15, 1.85, 2.25, 1.62],
        font_size=7.8,
    )

    add_heading(doc, "2.2 Scope Rules", 2)
    add_bullet(doc, "Division Chiefs see records assigned to committees under their Division.")
    add_bullet(doc, "Secretariat users see records assigned to their committees. A Lead Committee Secretariat label identifies records led by their committee.")
    add_bullet(doc, "Division Staff see records within the Division/committee scope assigned to their account.")
    add_bullet(doc, "Receiving Section Staff may see current status and tracking data but cannot change committee processing statuses.")
    add_bullet(doc, "Administrator access includes every role-specific feature, even when the Administrator dashboard is organized differently.")

    add_heading(doc, "2.3 Division Assignment", 2)
    add_body(doc, "Every account except Administrator and City Secretary must be assigned to one of these Divisions:")
    add_bullet(doc, "Legislative Committees Division")
    add_bullet(doc, "Legislative Support Services Division")
    add_bullet(doc, "Administrative Support Division")
    add_callout(
        doc,
        "Why assignment matters",
        "A missing or incorrect Division can hide records, committees, staff, dashboard counters, and reports. Ask an Administrator or City Secretary to correct the account.",
        "warning",
    )


def build_navigation(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 3",
        "Navigation and Dashboards",
        "The side menu and dashboard counters are the fastest way to find work that requires attention.",
    )
    add_heading(doc, "3.1 Main Menu", 2)
    add_table(
        doc,
        ["Menu", "Purpose", "Who commonly sees it"],
        [
            ("Dashboard", "Role-specific queues, counters, recent updates, search, committees, and staff.", "All signed-in users"),
            ("Records", "Search, filter, view, review, update, print, attach files, and manage permitted records.", "All signed-in users"),
            ("New Record", "Create a Committee Referral or Transmittals, Letters and Endorsements record.", "Administrator, City Secretary, Receiving Section Staff"),
            ("Management", "Committees, Terms, Officials Assignment, Secretariat Assignments, Committee Assignment, and User Creation.", "Authorized management roles"),
            ("Audit Logs", "Review user actions, date/time, item, description, and IP address.", "Administrator"),
            ("Backup", "Download a full SQL backup of schema and data.", "Administrator"),
            ("Reports", "Generate date-filtered status and committee summaries.", "All signed-in users, scoped by role"),
        ],
        [1.2, 3.65, 2.02],
        font_size=8.3,
    )

    add_heading(doc, "3.2 Attention Counters", 2)
    add_body(doc, "Red dashboard counters show records that need action. Open the corresponding tab to see the items. Records needing action are placed before ordinary records, with the newest action date first.")
    add_kpi_strip(
        doc,
        [
            ("RED COUNTER", "Action or review required"),
            ("NEW", "Updated today"),
            ("2 DAYS", "Orange age marker"),
            ("3+ DAYS", "Yellow age marker"),
        ],
        fill=GRAY,
    )
    add_body(doc, "The Records side-menu counter has been removed. Use dashboard tab counters as the authoritative workload indicator.")

    add_heading(doc, "3.3 Division Chief Dashboard", 2)
    add_bullet(doc, "Newly Assigned Referrals: records reviewed by the City Secretary and waiting for the Division Chief's first action.")
    add_bullet(doc, "Latest Staff Updates: latest Secretariat updates, Review Needed labels, comments, search, date filters, and pagination.")
    add_bullet(doc, "Committee: assigned committees, member rosters, and record counters.")
    add_bullet(doc, "Staff: full names, nicknames, login activity, and pending referral counters.")
    add_bullet(doc, "Search Record: searches assigned records.")

    add_heading(doc, "3.4 Secretariat Dashboard", 2)
    add_body(doc, "The Secretariat dashboard follows the same structure but is restricted to assigned committees. Latest Staff Updates is named Latest Updated Records. Newly assigned records become editable only after the Division Chief records the first action.")

    add_heading(doc, "3.5 City Secretary Dashboard", 2)
    add_bullet(doc, "Review: incoming records waiting for final review and classification.")
    add_bullet(doc, "Transmittals, Letters and Endorsements: the separate administrative record type.")
    add_bullet(doc, "For Plenary: records ready for proposed numbering and session processing.")
    add_bullet(doc, "Approved in the Plenary: approved records, retained even after LMIS & Records Staff adds later updates.")
    add_bullet(doc, "Recently Updated, Logs, Committees, and Search Record: system-wide monitoring tools.")


def build_record_intake(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 4",
        "Create a New Record",
        "Receiving Section Staff normally performs intake. Administrator and City Secretary may also create records.",
    )
    add_heading(doc, "4.1 Start a Record", 2)
    add_number(doc, "Select New Record from the side menu.")
    add_number(doc, "Choose the Record Type: Committee Referral or Transmittals, Letters and Endorsements.")
    add_number(doc, "Enter the Title. The field accepts long titles up to 1,000 characters.")
    add_number(doc, "Select Tag as New or Existing Record.")
    add_number(doc, "Enter the Client / Origin, Contact Number, and Email when available.")
    add_number(doc, "Select committee information when the document is expected to become a Committee Referral.")
    add_number(doc, "Confirm Date Received and Receiving Section Staff. These identify the intake event.")
    add_number(doc, "Attach supporting PDF or image files when available.")
    add_number(doc, "Select Save Record.")

    add_heading(doc, "4.2 Automatic Communication Number", 2)
    add_kpi_strip(
        doc,
        [
            ("COMMITTEE REFERRAL", "L-00001-YYYY"),
            ("TLE DOCUMENT", "A-00001-YYYY"),
            ("YEAR", "Changes automatically"),
        ],
    )
    add_body(doc, "The number is generated by the system. If the City Secretary changes the record type during review, the number changes to the next sequence for the selected type.")

    add_heading(doc, "4.3 New or Existing Record", 2)
    add_bullet(doc, "Tag as New creates an independent record.")
    add_bullet(doc, "Tag as Existing / Update Old Record lets the City Secretary merge the new intake into an existing record.")
    add_bullet(doc, "After merging, the communication uses the existing Communication Number and shares the history as one record.")
    add_callout(
        doc,
        "Before tagging as existing",
        "Search the old Communication Number and verify that the client/origin and subject refer to the same matter. An incorrect merge can confuse the permanent history.",
        "warning",
    )

    add_heading(doc, "4.4 Committee Selection and Lead Committee", 2)
    add_bullet(doc, "Select committees using the checkboxes in the committee menu.")
    add_bullet(doc, "With one selected committee, that committee automatically becomes the Lead Committee.")
    add_bullet(doc, "With two or more selected committees, select one Lead Committee from the selected choices.")
    add_bullet(doc, "The Lead Committee is shown first in record views and print referrals.")
    add_bullet(doc, "One Communication Number may be referred to 2 to 30 committees and remains one record with one history.")

    add_heading(doc, "4.5 Attachments at Intake", 2)
    add_bullet(doc, "Accepted file types: PDF and common image formats.")
    add_bullet(doc, "Maximum files on initial creation: 10.")
    add_bullet(doc, "Maximum per file: 10 MB. Maximum combined upload: 35 MB.")
    add_bullet(doc, "Select View All Documents to review several files as one document set.")


def build_committee_workflow(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 5",
        "Committee Referral Workflow",
        "The controlled handoff ensures intake, review, printing, committee action, and history are attributed to the correct role.",
    )
    add_workflow(
        doc,
        [
            ("1", "Receiving Section", "Create and attach"),
            ("2", "City Secretary", "Review and assign"),
            ("3", "Receiving Section", "Print and forward"),
            ("4", "Division Chief", "Record first action"),
            ("5", "Secretariat", "Process and update"),
            ("6", "City Secretary", "Plenary action"),
        ],
    )

    add_heading(doc, "5.1 Intake: Receiving Section Staff", 2)
    add_body(doc, "A new Committee Referral begins with status Received. Before City Secretary review, the creator may edit or delete the record. Once the record is reviewed and assigned, the intake data is no longer editable by Receiving Section Staff.")

    add_heading(doc, "5.2 Review: City Secretary", 2)
    add_number(doc, "Open the Review tab or select Review in Records.")
    add_number(doc, "Confirm or correct the record type, title, client/origin, and intake details.")
    add_number(doc, "Choose one or more committees. Confirm the Lead Committee.")
    add_number(doc, "Add review remarks when needed.")
    add_number(doc, "Save the review. The system displays the assigned Division Chief and Secretariat.")
    add_body(doc, "After review, the referral becomes Pending to the Committee and appears in the appropriate Division Chief and Receiving Section queues.")

    add_heading(doc, "5.3 Printing: Receiving Section Staff", 2)
    add_number(doc, "Open Records and locate the assigned referral.")
    add_number(doc, "Select View Record, then View Referral.")
    add_number(doc, "Review the folio preview and select Print.")
    add_number(doc, "After a successful print, return to the record and select Mark Printed and Forwarded.")
    add_number(doc, "Use Print Again whenever another copy is required.")
    add_callout(
        doc,
        "Printing indicator",
        "The For Printing action and red attention indicator remain until Mark Printed and Forwarded is recorded.",
        "important",
    )

    add_heading(doc, "5.4 First Action: Division Chief", 2)
    add_body(doc, "The Division Chief opens Newly Assigned Referrals and selects Action. The popup identifies the assigned Secretariat. Saving the first action records it in tracking history and enables Secretariat updates.")

    add_heading(doc, "5.5 Secretariat Processing", 2)
    add_body(doc, "After the first Division Chief action, the Secretariat may update status, enter action-specific details, upload attachments, and print a referral for the selected movement. The latest Secretariat update remains editable until another update is added so the Division Chief can review it.")

    add_heading(doc, "5.6 Division Chief Review", 2)
    add_body(doc, "Latest Staff Updates places Review Needed records first. The Division Chief may open the record in a popup, review the full history, edit the latest permitted Secretariat update, and add comments or instructions. Chief comments are internal guidance and do not print as referral remarks.")


def build_status_updates(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 6",
        "Update Status and History",
        "Every saved update creates a dated movement. Select the status that describes the action actually taken.",
    )
    add_heading(doc, "6.1 Save a Status Update", 2)
    add_number(doc, "Open the record and select Update Status, or use Action from the dashboard.")
    add_number(doc, "Choose a status from the list allowed for your role and the current record stage.")
    add_number(doc, "Complete the additional field shown for that status.")
    add_number(doc, "Enter Remarks when useful. Remarks are optional for all updates.")
    add_number(doc, "Attach supporting PDF/image files when required.")
    add_number(doc, "Select Save Update.")
    add_body(doc, "The update timestamp is created automatically and appears in Status Tracking History. The separate action date is used only when the committee or plenary action itself requires a date.")

    add_heading(doc, "6.2 Committee Status Reference", 2)
    add_table(
        doc,
        ["Status", "Additional input", "Operating note"],
        [
            ("Received", "Receiving Section Staff name", "Created automatically at intake."),
            ("Assigned to the Committee", "Committee", "System assignment event; not an ordinary Secretariat/Chief choice."),
            ("Pending to the Committee", "Committee", "Current committee stage after City Secretary review."),
            ("For Meeting", "Date of Meeting", "Use the scheduled/actual meeting date required by office procedure."),
            ("For Inspection", "Date of Inspection", "Use the inspection date."),
            ("Recommending Approval", "Recommended Committee and Date", "Adds the new committee to the same Communication Number and shared history."),
            ("Deferred", "Date Deferred", "Records deferral for later action."),
            ("Tabled", "Date Tabled", "Records the table action."),
            ("Noted", "Date Noted", "Records the note action."),
        ],
        [1.45, 1.65, 3.77],
        font_size=8.3,
    )

    add_heading(doc, "6.3 Remaining Committee Statuses", 2)
    add_table(
        doc,
        ["Status", "Additional input", "Operating note"],
        [
            ("Referred To", "Office / Organization / Individual", "The destination is shown in authorized and public status views."),
            ("Referred Back to Committee", "Committee dropdown", "Returns the matter to a selected committee."),
            ("Perusal", "No required special field", "The word Perusal appears on the upper-left portion of the Committee Report print."),
            ("Endorsement", "Date", "Records the endorsement date."),
            ("For Plenary Session", "Plenary / Session Date", "Moves the record to the City Secretary For Plenary tab."),
            ("Disapproved", "Date of Session", "Records plenary/committee disapproval; remarks remain optional."),
            ("Approved in the Plenary", "Type, Number, Date Approved", "City Secretary/Admin only; locks ordinary committee editing."),
            ("Others", "Details", "Use only when no listed status accurately describes the action."),
        ],
        [1.45, 1.75, 3.67],
        font_size=8.3,
    )

    add_heading(doc, "6.4 Titles and Remarks on Update Prints", 2)
    add_bullet(doc, "A Secretariat may enter an Updated Referral Title for a movement. The original record title remains unchanged.")
    add_bullet(doc, "Committee referral printouts use the movement title when one was entered.")
    add_bullet(doc, "Remarks are recorded in history but are no longer printed inside the Committee Report box.")
    add_bullet(doc, "Print Referral links may appear beside eligible movements for authorized roles.")

    add_heading(doc, "6.5 Status Tracking History", 2)
    add_body(doc, "The history shows the account nickname, date and time in MM/DD/YYYY Time format, the update/status, and all relevant data entered for that movement. The Administrator audit log also records broader user actions, including print activity.")
    add_callout(
        doc,
        "Do not rewrite history",
        "Correct only the latest update when the system permits it. Add a new update when the record has moved forward. This preserves the chronological audit trail.",
        "important",
    )


def build_plenary(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 7",
        "Plenary and Post-Plenary Processing",
        "The City Secretary manages proposed and approved numbers. LMIS & Records Staff manages the approved record after plenary.",
    )
    add_heading(doc, "7.1 For Plenary", 2)
    add_body(doc, "A record with status For Plenary Session appears in the City Secretary For Plenary tab. The City Secretary may filter records by plenary/session date and print the filtered result.")
    add_number(doc, "Open For Plenary and locate the record.")
    add_number(doc, "Select Assign Proposed No.")
    add_number(doc, "Choose Proposed Ordinance or Proposed Resolution.")
    add_number(doc, "Enter the proposed number. The system prefixes the current year, for example 2026-232.")
    add_number(doc, "Confirm the Plenary / Session Date and save.")
    add_body(doc, "The proposed type/number and session date are displayed together at the bottom of the record summary.")

    add_heading(doc, "7.2 Approve in the Plenary", 2)
    add_number(doc, "Open the plenary record and select Update Status.")
    add_number(doc, "Choose Approved in the Plenary.")
    add_number(doc, "Choose Ordinance or Resolution.")
    add_number(doc, "Enter the official number and Date Approved.")
    add_number(doc, "Save once. The system formats the identifier as Number-Year.")
    add_callout(
        doc,
        "Permanent lock",
        "After approval, Division Chief, Secretariat, and Receiving Section Staff cannot edit the record. Approved in the Plenary cannot be applied twice. Date Approved is shown as MM/DD/YYYY.",
        "warning",
    )

    add_heading(doc, "7.3 Approved in the Plenary Tab", 2)
    add_bullet(doc, "Filter by Date Approved.")
    add_bullet(doc, "Filter by Ordinance or Resolution.")
    add_bullet(doc, "The record remains in this City Secretary tab even after later LMIS & Records Staff updates.")
    add_bullet(doc, "Open View Record to see the full history and post-plenary status.")

    add_heading(doc, "7.4 LMIS & Records Staff Workflow", 2)
    add_body(doc, "LMIS & Records Staff works only with approved/post-plenary records in its assigned dashboard. Each update requires an action date; remarks are optional.")
    add_table(
        doc,
        ["Post-plenary status", "Typical use"],
        [
            ("For Vice Mayor's Signature", "Sent for the Vice Mayor's signature."),
            ("Returned from The Vice Mayor", "Returned after Vice Mayor action."),
            ("Forwarded for Admin/Mayor Signature", "Sent to Administration/Mayor for signature."),
            ("Returned from Admin/Mayor", "Returned after Administration/Mayor action."),
            ("Veto", "Records veto action."),
            ("Lapse into Ordinance", "Records effect by lapse into ordinance."),
            ("Forwarded to the Messengerial Services", "Sent for delivery/service."),
            ("For Transmittal", "Ready for recipient management and transmittal."),
            ("Completed", "Post-plenary processing is complete."),
        ],
        [2.85, 4.02],
        font_size=8.4,
    )

    add_heading(doc, "7.5 Add Transmittal Recipients", 2)
    add_number(doc, "Open an approved record with status For Transmittal.")
    add_number(doc, "Select Add Recipients.")
    add_number(doc, "Enter Title, Name, Position, Address, and Contact Number.")
    add_number(doc, "Save each recipient and review the list before delivery.")
    add_callout(
        doc,
        "Terminology",
        "The post-plenary Transmittals tab is not the same as the record type Transmittals, Letters and Endorsements. The former contains approved legislative records ready for service.",
        "important",
    )


def build_admin_docs(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 8",
        "Transmittals, Letters and Endorsements",
        "This record type is separate from Committee Referrals and from post-plenary transmittal processing.",
    )
    add_heading(doc, "8.1 Create the Record", 2)
    add_number(doc, "Open New Record.")
    add_number(doc, "Choose Transmittals, Letters and Endorsements.")
    add_number(doc, "Enter the Title and Client / Origin.")
    add_number(doc, "Enter the Contact Number and Email when available.")
    add_number(doc, "Confirm Date Received and Receiving Section Staff.")
    add_number(doc, "Attach supporting files and save.")
    add_body(doc, "The system assigns an A-series Communication Number in the form A-00001-YYYY.")

    add_heading(doc, "8.2 City Secretary Review and Forwarding", 2)
    add_number(doc, "Open the Transmittals, Letters and Endorsements tab.")
    add_number(doc, "Select Review.")
    add_number(doc, "Choose Forwarded To: City Vice Mayor, City Councilors, Divisions, Sections, or Employee.")
    add_number(doc, "Enter the destination details and remarks. The remarks field is emphasized after a forwarding choice.")
    add_number(doc, "Save the review.")

    add_heading(doc, "8.3 Status and Actions", 2)
    add_table(
        doc,
        ["Status", "Meaning"],
        [
            ("Received", "The record has been accepted by the Receiving Section."),
            ("Completed", "The requested administrative action is complete."),
            ("Archived", "The record is retained with no further active action."),
        ],
        [1.7, 5.17],
    )
    add_body(doc, "City Secretary actions for this record type are Review and Delete. Committee assignment, committee status updating, and referral printing do not apply.")

    add_callout(
        doc,
        "Client / Origin",
        "The Client field used for Committee Referrals and the Origin field used for Transmittals, Letters and Endorsements share the same stored information. Changing type during review must preserve this value.",
        "info",
    )


def build_records(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 9",
        "Records, Search, Attachments, and History",
        "Use Records for detailed work and dashboards for prioritized queues.",
    )
    add_heading(doc, "9.1 Record List", 2)
    add_body(doc, "Records are displayed with Communication Number, Type, Date Received, Client / Origin, Committee, Title, Division, Assigned Secretariat, current Status, attachments, and Actions. The latest or action-required records appear first according to the active view.")
    add_bullet(doc, "Review: City Secretary review/classification.")
    add_bullet(doc, "Update: permitted status update.")
    add_bullet(doc, "View Record: full record details and history.")
    add_bullet(doc, "Print / For Printing: Receiving Section referral printing workflow.")
    add_bullet(doc, "Delete: Administrator, City Secretary, or qualifying unassigned creator.")

    add_heading(doc, "9.2 Search and Filters", 2)
    add_number(doc, "Enter any known word or number: Communication Number, title, client/origin, status, remarks, or staff name.")
    add_number(doc, "Set From Date and To Date when narrowing by period.")
    add_number(doc, "Use committee, status, type, or dashboard-specific filters when available.")
    add_number(doc, "Select Search.")
    add_number(doc, "Select Clear to remove filters.")
    add_body(doc, "When opening a record from a dashboard popup, Close returns to the same tab and preserves search, date, and page settings.")

    add_heading(doc, "9.3 View a Record", 2)
    add_bullet(doc, "Review the current status and complete record details.")
    add_bullet(doc, "For multi-committee records, confirm that all committees, Division Chiefs, and Secretariats are shown, with the Lead Committee first.")
    add_bullet(doc, "Check proposed or approved numbers and dates when applicable.")
    add_bullet(doc, "Use Status Tracking History for the full chronology.")
    add_bullet(doc, "Use the available buttons for updates, comments, printing, attachments, and plenary actions.")

    add_heading(doc, "9.4 Multiple Attachments", 2)
    add_body(doc, "Authorized Receiving Section Staff, Secretariat, Division Chief, and Administrator users may upload PDF/image attachments according to record scope. Individual attachment links are consolidated into View All Documents in the record view.")
    add_callout(
        doc,
        "Sensitive documents",
        "Check that the correct files are attached before saving. Attachments are available to authorized users who can view the record and may contain personal or official information.",
        "warning",
    )

    add_heading(doc, "9.5 Public Request Status", 2)
    add_body(doc, "People without accounts can select Check the Status of your Request on the public main window, enter a Communication Number, and view the current public status. The page shows the latest status and a Referred To destination when applicable, but not the full internal history.")


def build_printing(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 10",
        "Committee Referral Printing",
        "The print template is designed as a formal folio document and uses committee, term, official, and movement data.",
    )
    add_heading(doc, "10.1 Document Contents", 2)
    add_bullet(doc, "Official SP logo and government office header.")
    add_bullet(doc, "Communication Number, date, client/origin, and full title.")
    add_bullet(doc, "Referral paragraph naming the committee.")
    add_bullet(doc, "City Secretary signature block on the initial referral. The signature image is omitted on Secretariat-generated update prints.")
    add_bullet(doc, "Committee Report, committee code-based Report Number, action date, and report writing box.")
    add_bullet(doc, "Chairperson, Vice Chairperson, committee members, ex-officio members, and dissenting signature spaces.")
    add_bullet(doc, "Joint/Tripartite/etc. sequence note for multi-committee referrals, with the Lead Committee first.")
    add_bullet(doc, "Generated timestamp tied to the original creation/printing event.")

    add_heading(doc, "10.2 Print Settings", 2)
    add_kpi_strip(
        doc,
        [
            ("PAPER", "Folio 8.5 x 13"),
            ("ORIENTATION", "Portrait"),
            ("SCALE", "100%"),
            ("HEADERS/FOOTERS", "Browser off"),
        ],
    )
    add_number(doc, "Open View Referral or Print Referral.")
    add_number(doc, "Confirm the on-screen folio preview.")
    add_number(doc, "Select Print. In Chrome or Edge, choose the correct printer.")
    add_number(doc, "Set paper to Folio / 8.5 x 13 inches and scale to 100%.")
    add_number(doc, "Disable browser headers and footers. Use the page margins supplied by the template.")
    add_number(doc, "Print one test copy before producing multiple official copies.")

    add_heading(doc, "10.3 Print Does Not Open", 2)
    add_bullet(doc, "Use Ctrl+P while the referral preview is active.")
    add_bullet(doc, "Allow popups/printing for the local system address.")
    add_bullet(doc, "Try current Chrome or Microsoft Edge.")
    add_bullet(doc, "Confirm a printer is installed and Windows can print a test page.")

    add_heading(doc, "10.4 Mark Printed and Forwarded", 2)
    add_body(doc, "Only use Mark Printed and Forwarded after the official referral was successfully produced and forwarded. This removes the Receiving Section attention indicator while preserving Print Again.")


def build_management(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 11",
        "Terms, Officials, Committees, and Assignments",
        "Accurate reference data controls committee rosters, signatures, report numbers, and role access.",
    )
    add_heading(doc, "11.1 Terms", 2)
    add_number(doc, "Open Management > Terms.")
    add_number(doc, "Enter the Term Name, Start Year, and End Year.")
    add_number(doc, "Select Current when this is the active three-year term.")
    add_number(doc, "Save. Use Edit or Delete only when authorized and no dependent data will be damaged.")
    add_callout(
        doc,
        "Term changes",
        "Create a new term for a new council period. Do not overwrite an old term, because historical rosters must remain connected to earlier records.",
        "important",
    )

    add_heading(doc, "11.2 Officials Assignment", 2)
    add_number(doc, "Open Management > Officials Assignment and select a term.")
    add_number(doc, "Enter the official's name and elected Position.")
    add_number(doc, "Select an Officer Role when applicable.")
    add_number(doc, "Set Display Order and save.")
    add_body(doc, "Officer Roles include Presiding Officer, Presiding Officer Pro-Tempore, Majority/Minority Floor Leaders and assistants, SK Federation President, IPMR Representative, and Association of Barangay Captain President.")

    add_heading(doc, "11.3 Committees and Rosters", 2)
    add_number(doc, "Open Management > Committees.")
    add_number(doc, "Enter the Committee Name, Committee Code, and Description.")
    add_number(doc, "Choose the term and assign Chairperson, Vice Chairperson, and up to 20 Members.")
    add_number(doc, "Save. Officials are displayed alphabetically where applicable.")
    add_body(doc, "The Committee Code generates Committee Report numbers in the format CODE-0001-YY. The See Members button opens the current roster in a popup.")

    add_heading(doc, "11.4 Committee Assignment", 2)
    add_body(doc, "Administrator or City Secretary assigns each committee to one Division Chief. A committee cannot be assigned to two Division Chiefs; the system displays an error if a duplicate assignment is attempted.")

    add_heading(doc, "11.5 Secretariat Assignments", 2)
    add_body(doc, "Assign Secretariat users to the committees they support. A Division Chief sees only staff and committees within their scope. These assignments directly control dashboards, record access, counters, and update permissions.")


def build_users(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 12",
        "User Accounts",
        "Administrator and City Secretary create and maintain accounts. Division fields determine record visibility.",
    )
    add_heading(doc, "12.1 Create an Account", 2)
    add_number(doc, "Open Management > User Creation.")
    add_number(doc, "Select Create User.")
    add_number(doc, "Enter Full Name and Nickname.")
    add_number(doc, "Choose the Division. Administrator and City Secretary are exempt.")
    add_number(doc, "Enter a unique Email address.")
    add_number(doc, "Choose the Role.")
    add_number(doc, "Enter an initial password.")
    add_number(doc, "Select Check to Activate Account beside Save.")
    add_number(doc, "Select Save.")
    add_body(doc, "User lists are grouped by Division. Administrator and City Secretary accounts are displayed separately.")

    add_heading(doc, "12.2 Select the Correct Role", 2)
    add_table(
        doc,
        ["Role choice", "Choose when the person will..."],
        [
            ("Administrator", "Manage the entire system, logs, backup, and all functions."),
            ("City Secretary", "Review all records, manage committees/users, and process plenary actions."),
            ("Division Chief", "Lead assigned committees and review Secretariat work."),
            ("Receiving Section Staff", "Receive and create records, upload intake files, and print initial referrals."),
            ("Secretariat", "Process records for assigned committees."),
            ("Division Staff", "Search and view records within a Division without editing."),
            ("LMIS & Records Staff", "Process approved records and post-plenary transmittals."),
        ],
        [2.0, 4.87],
    )

    add_heading(doc, "12.3 Edit, Activate, and Reset Password", 2)
    add_number(doc, "Select Edit beside the account. The form opens in a popup.")
    add_number(doc, "Correct the name, nickname, division, email, role, or activation state.")
    add_number(doc, "Enter a new password only when resetting it.")
    add_number(doc, "Save and close the popup.")
    add_callout(
        doc,
        "Passwords cannot be viewed",
        "The system stores password hashes, not readable passwords. Even an Administrator cannot retrieve the current password. Reset it to a new temporary password and tell the user securely.",
        "warning",
    )

    add_heading(doc, "12.4 Account Maintenance", 2)
    add_bullet(doc, "Deactivate accounts promptly when a person transfers or leaves.")
    add_bullet(doc, "Update Division and committee assignment when responsibilities change.")
    add_bullet(doc, "Review Last Login information when investigating access issues.")
    add_bullet(doc, "Never reuse a shared office account for several employees.")


def build_reports_admin(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 13",
        "Reports, Audit Logs, and Backup",
        "Reports support operations. Audit and backup functions protect accountability and continuity.",
    )
    add_heading(doc, "13.1 Generate a Report", 2)
    add_number(doc, "Open Reports.")
    add_number(doc, "Set Date From and Date To.")
    add_number(doc, "Select Generate Report.")
    add_number(doc, "Review totals By Status, By Committee, and Recent Records.")
    add_number(doc, "Select Print to produce the report.")
    add_body(doc, "Division Chief, Secretariat, and Division Staff reports are restricted to their permitted committee scope. Administrator and City Secretary can review broader system totals.")

    add_heading(doc, "13.2 Audit Logs", 2)
    add_body(doc, "Administrator only. Open Audit Logs and filter by User, Action, and Date. Each entry may show date/time, user, action, affected item, description, and IP address. Use logs to investigate changes, printing, account activity, and administrative actions.")
    add_callout(
        doc,
        "Tracking history vs. audit log",
        "Status Tracking History explains the record's official movement. Audit Logs show broader system activity and are the correct place to review all print and user actions.",
        "info",
    )

    add_heading(doc, "13.3 Download a Backup", 2)
    add_number(doc, "Sign in as Administrator.")
    add_number(doc, "Open Backup.")
    add_number(doc, "Select Download Backup.")
    add_number(doc, "Store the downloaded .sql file in a protected location separate from the host computer.")
    add_body(doc, "The file name follows lcd_records_backup_YYYYMMDD_HHMMSS.sql and contains the database schema and data.")

    add_heading(doc, "13.4 Suggested Backup Routine", 2)
    add_table(
        doc,
        ["Frequency", "Action"],
        [
            ("Daily", "Download a database backup after the last major update period."),
            ("Weekly", "Copy database backup and uploaded attachments to protected external/network storage."),
            ("Before upgrades", "Back up database, public assets, application files, and storage/attachments."),
            ("Monthly", "Test restoration on a separate non-production database."),
        ],
        [1.35, 5.52],
    )

    add_heading(doc, "13.5 Restore Overview", 2)
    add_body(doc, "Restoration should be performed by the Administrator or IT support. Create/select the target MySQL database, import the .sql backup in phpMyAdmin or the MySQL command line, restore the matching attachments folder, verify app/config.php, and test sign-in and records before reopening access.")
    add_callout(
        doc,
        "Never overwrite without a backup",
        "A restore can replace current data. Make a fresh backup and confirm the target database before importing.",
        "warning",
    )


def build_network(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 14",
        "Running on a Local Network",
        "The host computer runs PHP and MySQL; other computers connect through the host's network address.",
    )
    add_heading(doc, "14.1 Host Computer Checklist", 2)
    add_number(doc, "Connect the host and user computers to the same trusted network.")
    add_number(doc, "Start MySQL in XAMPP or the configured database service.")
    add_number(doc, "Start the application web server so it listens on all network interfaces.")
    p = doc.add_paragraph()
    style_paragraph(p, space_after=7, line=1)
    set_run_font(p.add_run(r"C:\dbase\php\php.exe -S 0.0.0.0:8080 -t public"), size=9, bold=True, color=NAVY, name="Consolas")
    add_number(doc, "Allow PHP or TCP port 8080 through Windows Firewall on the Private network.")
    add_number(doc, "Find the host IPv4 address using ipconfig.")
    add_number(doc, "On another computer, open http://HOST-IP:8080/.")

    add_heading(doc, "14.2 Example", 2)
    add_body(doc, "If the host IPv4 address is 192.168.1.25, users open:")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell = doc.add_table(rows=1, cols=1)
    set_table_grid(set_cell, [4.6])
    set_cell.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = set_cell.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_margins(cell, 130, 140, 130, 140)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cp.add_run("http://192.168.1.25:8080/"), size=13, bold=True, color=BLUE, name="Consolas")

    add_heading(doc, "14.3 Operating Rules", 2)
    add_bullet(doc, "The host computer, PHP server, and MySQL must remain on while users are connected.")
    add_bullet(doc, "Use one central application folder, database, and storage/attachments folder.")
    add_bullet(doc, "Give the host computer a reserved/static local IP so the address does not frequently change.")
    add_bullet(doc, "Use HTTPS, proper authentication, and IT-managed hosting before exposing the system outside the office network.")
    add_bullet(doc, "Do not forward MySQL directly to user computers; users connect through the web application.")

    add_heading(doc, "14.4 XAMPP Apache Alternative", 2)
    add_body(doc, "When deployed under XAMPP Apache, place the application in the configured web directory and browse using the host IP plus the application path. Confirm that Apache's document root points to the public folder or is configured so app and database files are not publicly served.")


def build_troubleshooting(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 15",
        "Troubleshooting",
        "Use these checks before escalating. Record the exact error message, page address, user role, and time.",
    )
    add_table(
        doc,
        ["Problem", "Likely cause", "Action"],
        [
            ("Site does not open", "PHP/Apache is stopped or wrong port/address.", "Start the web server. On another PC, use the host IP, not 127.0.0.1."),
            ("MySQL shutdown unexpectedly", "Port conflict, crash, permissions, or damaged data files.", "Check the XAMPP MySQL log and Windows Event Viewer. Confirm only one MySQL service uses the port. Restore only with a verified backup."),
            ("Sign-in loops or stays on login", "Database unavailable, account inactive, or session storage issue.", "Start MySQL, verify credentials/account activation, refresh, then check storage/sessions is writable."),
            ("session_start permission denied", "PHP cannot write the session directory.", "Use the project storage/sessions directory and grant the web server account write permission."),
            ("Updates not seen on another PC", "Users opened different local copies/databases.", "Make every user open the same host IP and confirm one central database."),
            ("Record is missing", "Role, Division, or committee scope excludes it.", "Check account Division and committee assignment with Administrator/City Secretary."),
            ("Secretariat cannot update", "Division Chief has not recorded the first action.", "Chief opens Newly Assigned Referrals and saves Action."),
            ("No print option", "Record is unassigned, already handled, wrong role, or popup blocked.", "Confirm City Secretary review, role, For Printing state, and popup permissions."),
            ("Referral spills to page 2", "Wrong paper/scale/browser headers.", "Choose Folio 8.5 x 13, 100% scale, portrait, and disable browser headers/footers."),
            ("Attachments unavailable", "Different server copy, missing storage folder, size/type failure.", "Use central host, verify storage/attachments, and check PDF/image limits."),
            ("Unknown database column", "A migration was not applied.", "Back up first, then apply the matching migration once or install the current schema on a fresh database."),
            ("Password forgotten", "Passwords are non-recoverable hashes.", "Administrator/City Secretary resets the password; do not attempt to display it."),
        ],
        [1.35, 2.08, 3.44],
        font_size=7.75,
    )

    add_heading(doc, "15.1 Information to Give IT Support", 2)
    add_bullet(doc, "Exact page address and full error message.")
    add_bullet(doc, "User role, Division, committee, and Communication Number.")
    add_bullet(doc, "Date and time the issue occurred.")
    add_bullet(doc, "Whether the issue occurs on the host computer, another computer, or both.")
    add_bullet(doc, "Whether PHP/Apache and MySQL are running.")
    add_bullet(doc, "A screenshot that does not expose passwords or confidential data.")

    add_callout(
        doc,
        "Protect evidence",
        "Do not delete database files, reset MySQL data directories, reinstall XAMPP, or overwrite the application before a current backup is secured.",
        "warning",
    )


def build_quick_reference(doc: Document) -> None:
    add_section_page(
        doc,
        "Section 16",
        "Quick Reference",
        "Daily checklists and escalation points for consistent operation.",
    )
    add_heading(doc, "16.1 Receiving Section Staff", 2)
    add_bullet(doc, "Check new communications and create complete intake records.")
    add_bullet(doc, "Attach the received PDF/images and verify Client / Origin contact data.")
    add_bullet(doc, "Correct or delete your own record only while it remains Received and unassigned.")
    add_bullet(doc, "Monitor For Printing, print assigned referrals, and mark Printed and Forwarded.")
    add_bullet(doc, "Use View Record and history to answer status inquiries.")

    add_heading(doc, "16.2 City Secretary", 2)
    add_bullet(doc, "Open Review and clear action-required records.")
    add_bullet(doc, "Confirm type, committee(s), Lead Committee, Division, and Secretariat.")
    add_bullet(doc, "Monitor Transmittals, Letters and Endorsements and recent updates.")
    add_bullet(doc, "Process For Plenary proposed numbers, filters, and session lists.")
    add_bullet(doc, "Apply Approved in the Plenary only once with correct type, number, and date.")

    add_heading(doc, "16.3 Division Chief", 2)
    add_bullet(doc, "Review Newly Assigned Referrals and record the first action.")
    add_bullet(doc, "Clear Review Needed items in Latest Staff Updates.")
    add_bullet(doc, "Use comments/instructions to guide the assigned Secretariat.")
    add_bullet(doc, "Check committee and staff counters for pending work.")

    add_heading(doc, "16.4 Secretariat", 2)
    add_bullet(doc, "Act only on records assigned to your committees.")
    add_bullet(doc, "Wait for the Division Chief first action before updating.")
    add_bullet(doc, "Enter the correct status-specific field and optional remarks.")
    add_bullet(doc, "Upload supporting files and confirm the latest movement.")

    doc.add_page_break()
    add_heading(doc, "16.5 LMIS & Records Staff", 2)
    add_bullet(doc, "Monitor approved/post-plenary records.")
    add_bullet(doc, "Record action dates and post-plenary statuses.")
    add_bullet(doc, "Add and verify recipients when status is For Transmittal.")
    add_bullet(doc, "Mark Completed only after all required processing is finished.")

    add_heading(doc, "16.6 Administrator", 2)
    add_bullet(doc, "Review account, Division, committee, and role assignments.")
    add_bullet(doc, "Review audit logs for unusual or disputed activity.")
    add_bullet(doc, "Download and protect regular backups, including attachments.")
    add_bullet(doc, "Confirm the host, PHP/Apache, MySQL, firewall, and network address.")

    add_heading(doc, "16.7 Support Contacts", 2)
    add_table(
        doc,
        ["Issue", "First contact"],
        [
            ("Record content, committee, or status question", "City Secretary or assigned Division Chief"),
            ("Account, role, Division, or permission issue", "Administrator or City Secretary"),
            ("Server, database, network, or backup issue", "Administrator / designated IT support"),
            ("Printing format or printer issue", "Receiving Section supervisor / IT support"),
        ],
        [2.65, 4.22],
    )

    add_callout(
        doc,
        "End of manual",
        "Keep this manual with the system's approved operating procedures. Update the version whenever roles, statuses, workflows, or menus change.",
        "success",
    )


def build_document() -> Path:
    doc = Document()
    format_document(doc)
    build_cover(doc)
    build_front_matter(doc)
    build_introduction(doc)
    build_roles(doc)
    build_navigation(doc)
    build_record_intake(doc)
    build_committee_workflow(doc)
    build_status_updates(doc)
    build_plenary(doc)
    build_admin_docs(doc)
    build_records(doc)
    build_printing(doc)
    build_management(doc)
    build_users(doc)
    build_reports_admin(doc)
    build_network(doc)
    build_troubleshooting(doc)
    build_quick_reference(doc)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
